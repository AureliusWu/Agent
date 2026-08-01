"""DOCX adapter backed by python-docx 1.2.0 and the shared OOXML layer."""

from __future__ import annotations

import io
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from app.artifacts.errors import ArtifactError
from app.artifacts.model import (
    ArtifactDocument,
    ArtifactTable,
    ContentBlock,
    TextReplacement,
    apply_text_replacements,
    normalize_document,
    normalize_replacements,
)
from app.artifacts.ooxml import (
    embed_model_metadata,
    inspect_ooxml,
    read_model_metadata,
    resolve_workspace_image,
)

_CJK_FONT = "Microsoft YaHei"


def _apply_run_font(run: Any, *, size_pt: float | None = None, bold: bool | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = _CJK_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _CJK_FONT)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def _configure_styles(document: Any) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    sizes = {
        "Normal": 11.0,
        "Title": 20.0,
        "Heading 1": 16.0,
        "Heading 2": 14.0,
        "Heading 3": 12.0,
    }
    for style_name, size in sizes.items():
        if style_name not in document.styles:
            continue
        style = document.styles[style_name]
        style.font.name = _CJK_FONT
        style.font.size = Pt(size)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _CJK_FONT)


def _add_table(document: Any, table_model: ArtifactTable) -> None:
    table = document.add_table(
        rows=len(table_model.rows),
        cols=len(table_model.rows[0]),
    )
    table.style = "Table Grid"
    for row_index, row in enumerate(table_model.rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _apply_run_font(
                        run,
                        bold=row_index < table_model.header_rows,
                    )


def _add_image(document: Any, block: ContentBlock, workspace_root: Path | None) -> None:
    from docx.shared import Inches

    assert block.image is not None
    path = resolve_workspace_image(block.image.source, workspace_root)
    width = None
    if block.image.width_px is not None:
        width = Inches(min(block.image.width_px / 96.0, 6.3))
    shape = document.add_picture(str(path), width=width)
    alt_text = block.image.alt_text.strip()
    if alt_text:
        shape._inline.docPr.set("descr", alt_text)
        shape._inline.docPr.set("title", alt_text[:255])


def _append_block(document: Any, block: ContentBlock, workspace_root: Path | None) -> None:
    if block.kind == "heading":
        paragraph = document.add_heading(block.text, level=block.level or 1)
        for run in paragraph.runs:
            _apply_run_font(run, bold=True)
    elif block.kind == "paragraph":
        paragraph = document.add_paragraph()
        _apply_run_font(paragraph.add_run(block.text))
    elif block.kind in {"bullet_list", "numbered_list"}:
        style = "List Bullet" if block.kind == "bullet_list" else "List Number"
        for item in block.items:
            paragraph = document.add_paragraph(style=style)
            _apply_run_font(paragraph.add_run(item))
    elif block.kind == "page_break":
        document.add_page_break()
    elif block.kind == "table":
        assert block.table is not None
        _add_table(document, block.table)
    elif block.kind == "image":
        _add_image(document, block, workspace_root)
    else:  # pragma: no cover - the normalized model rejects this first
        raise ArtifactError(
            "unsupported_artifact_block",
            "create",
            "The DOCX adapter received an unsupported block.",
            details={"kind": block.kind},
        )


def _assert_docx_reopens(package: bytes) -> None:
    from app.artifacts.validation import validate_docx

    report = validate_docx(package)
    if not report.passed:
        first = report.errors[0] if report.errors else None
        raise ArtifactError(
            first.error_code if first else "docx_reopen_failed",
            "validate",
            "The generated DOCX file could not be independently reopened.",
            details=first.details if first else {},
        )


def create_docx(
    document: ArtifactDocument | Mapping[str, Any],
    *,
    workspace_root: Path | None = None,
) -> bytes:
    """Create a bounded DOCX in memory; the service layer owns filesystem writes."""

    normalized = normalize_document(document)
    try:
        from docx import Document

        output_document = Document()
        _configure_styles(output_document)
        output_document.core_properties.title = normalized.title
        if normalized.title:
            paragraph = output_document.add_paragraph(style="Title")
            _apply_run_font(paragraph.add_run(normalized.title), size_pt=20, bold=True)
        for block in normalized.blocks:
            _append_block(output_document, block, workspace_root)
        buffer = io.BytesIO()
        output_document.save(buffer)
        package = embed_model_metadata(buffer.getvalue(), "docx", normalized)
        _assert_docx_reopens(package)
        return package
    except ArtifactError:
        raise
    except Exception as exc:
        raise ArtifactError(
            "docx_create_failed",
            "create",
            "The DOCX file could not be created.",
            details={"reason": type(exc).__name__},
        ) from None


def _iter_table_paragraphs(table: Any) -> Iterable[Any]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _iter_all_paragraphs(document: Any) -> Iterable[Any]:
    seen: set[int] = set()

    def emit(paragraphs: Iterable[Any]) -> Iterable[Any]:
        for paragraph in paragraphs:
            identity = id(paragraph._p)
            if identity not in seen:
                seen.add(identity)
                yield paragraph

    yield from emit(document.paragraphs)
    for table in document.tables:
        yield from emit(_iter_table_paragraphs(table))


def _replace_paragraph_text(paragraph: Any, old: str, new: str) -> int:
    original = "".join(run.text for run in paragraph.runs)
    count = original.count(old)
    if count == 0:
        return 0
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
    updated = "".join(run.text for run in paragraph.runs)
    if updated.count(old) == 0 and updated == original.replace(old, new):
        return count
    if not paragraph.runs:
        return 0
    paragraph.runs[0].text = original.replace(old, new)
    for run in paragraph.runs[1:]:
        run.text = ""
    return count


def _extract_table(table: Any) -> ArtifactTable:
    rows = tuple(
        tuple(cell.text for cell in row.cells)
        for row in table.rows
    )
    return ArtifactTable(rows=rows, header_rows=min(1, len(rows)))


def _extract_docx_without_metadata(document: Any) -> ArtifactDocument:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    blocks: list[ContentBlock] = []
    title = document.core_properties.title or ""
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text
            style_name = paragraph.style.name if paragraph.style is not None else ""
            if text:
                if style_name == "Title":
                    if not title:
                        title = text
                elif style_name.startswith("Heading "):
                    try:
                        level = max(1, min(6, int(style_name.rsplit(" ", 1)[1])))
                    except ValueError:
                        level = 1
                    blocks.append(ContentBlock(kind="heading", text=text, level=level))
                elif style_name.startswith("List Bullet"):
                    blocks.append(ContentBlock(kind="bullet_list", items=(text,)))
                elif style_name.startswith("List Number"):
                    blocks.append(ContentBlock(kind="numbered_list", items=(text,)))
                else:
                    blocks.append(ContentBlock(kind="paragraph", text=text))
            if child.xpath(".//*[local-name()='br' and @*[local-name()='type']='page']"):
                blocks.append(ContentBlock(kind="page_break"))
        elif child.tag.endswith("}tbl"):
            blocks.append(ContentBlock(kind="table", table=_extract_table(Table(child, document))))
    return ArtifactDocument(title=title, blocks=tuple(blocks))


def read_docx(package: bytes) -> ArtifactDocument:
    """Securely inspect, independently reopen, and normalize a DOCX."""

    inspect_ooxml(package, "docx")
    try:
        from docx import Document

        reopened = Document(io.BytesIO(package))
        metadata = read_model_metadata(package, "docx")
        return metadata if metadata is not None else _extract_docx_without_metadata(reopened)
    except ArtifactError:
        raise
    except Exception as exc:
        raise ArtifactError(
            "docx_read_failed",
            "read",
            "The DOCX file could not be read.",
            details={"reason": type(exc).__name__},
        ) from None


def edit_docx(
    package: bytes,
    replacements: (
        Mapping[str, str]
        | Sequence[TextReplacement | Mapping[str, Any]]
    ),
) -> bytes:
    """Apply count-bound exact visible-text replacements and return new bytes."""

    inspect_ooxml(package, "docx")
    normalized_replacements = normalize_replacements(replacements)
    try:
        from docx import Document

        existing_model = read_model_metadata(package, "docx")
        document = Document(io.BytesIO(package))
        for replacement in normalized_replacements:
            actual_count = sum(
                _replace_paragraph_text(
                    paragraph,
                    replacement.old,
                    replacement.new,
                )
                for paragraph in _iter_all_paragraphs(document)
            )
            if actual_count != replacement.expected_count:
                raise ArtifactError(
                    "artifact_replacement_count_mismatch",
                    "edit",
                    "Exact text replacement count did not match.",
                    details={
                        "expected_count": replacement.expected_count,
                        "actual_count": actual_count,
                    },
                )
        if existing_model is not None:
            updated_model, _ = apply_text_replacements(
                existing_model,
                normalized_replacements,
            )
            document.core_properties.title = updated_model.title
        buffer = io.BytesIO()
        document.save(buffer)
        raw_result = buffer.getvalue()
        if existing_model is None:
            reopened = Document(io.BytesIO(raw_result))
            updated_model = _extract_docx_without_metadata(reopened)
        result = embed_model_metadata(raw_result, "docx", updated_model)
        _assert_docx_reopens(result)
        return result
    except ArtifactError:
        raise
    except Exception as exc:
        raise ArtifactError(
            "docx_edit_failed",
            "edit",
            "The DOCX file could not be edited.",
            details={"reason": type(exc).__name__},
        ) from None


replace_docx_text = edit_docx
