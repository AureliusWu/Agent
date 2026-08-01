"""Unified independent-open validation results for document artifacts."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any, Literal

from app.artifacts.errors import ArtifactError, sanitize_details
from app.artifacts.ooxml import inspect_ooxml, read_model_metadata

ValidationStatus = Literal["PASS", "FAIL"]


@dataclass(frozen=True, slots=True)
class ValidationIssue:
    error_code: str
    phase: str
    message: str
    details: dict[str, Any] = field(default_factory=dict)

    @classmethod
    def from_error(cls, error: ArtifactError) -> "ValidationIssue":
        return cls(
            error_code=error.error_code,
            phase=error.phase,
            message=error.message,
            details=dict(error.safe_details),
        )

    def to_dict(self) -> dict[str, Any]:
        return {
            "error_code": self.error_code,
            "phase": self.phase,
            "message": self.message,
            "details": sanitize_details(self.details),
        }


@dataclass(frozen=True, slots=True)
class ValidationReport:
    format: str
    status: ValidationStatus
    openable: bool
    page_count: int | None = None
    slide_count: int | None = None
    checks: tuple[str, ...] = ()
    risks: tuple[ValidationIssue, ...] = ()
    errors: tuple[ValidationIssue, ...] = ()

    @property
    def count(self) -> int | None:
        return self.slide_count if self.slide_count is not None else self.page_count

    @property
    def passed(self) -> bool:
        return self.status == "PASS"

    def to_dict(self) -> dict[str, Any]:
        return {
            "status": self.status,
            "format": self.format,
            "openable": self.openable,
            "page_count": self.page_count,
            "slide_count": self.slide_count,
            "count": self.count,
            "checks": list(self.checks),
            "risks": [risk.to_dict() for risk in self.risks],
            "errors": [error.to_dict() for error in self.errors],
        }


def _failed_report(artifact_format: str, error: ArtifactError) -> ValidationReport:
    return ValidationReport(
        format=artifact_format,
        status="FAIL",
        openable=False,
        errors=(ValidationIssue.from_error(error),),
    )


def _engine_docx_page_count(package: bytes) -> int | None:
    model = read_model_metadata(package, "docx")
    if model is None:
        return None
    return 1 + sum(block.kind == "page_break" for block in model.blocks)


def validate_docx(package: bytes) -> ValidationReport:
    """Inspect and reopen DOCX through an independent public parser."""

    try:
        inspection = inspect_ooxml(package, "docx")
        from docx import Document

        reopened = Document(io.BytesIO(package))
        _ = tuple(paragraph.text for paragraph in reopened.paragraphs)
        page_count = _engine_docx_page_count(package)
        return ValidationReport(
            format="docx",
            status="PASS",
            openable=True,
            page_count=page_count,
            checks=(
                "bounded_ooxml_package",
                "content_types",
                "relationships",
                "macro_free",
                "external_relationship_free",
                "python_docx_reopen",
                *(
                    ("engine_metadata",)
                    if inspection.has_engine_metadata
                    else ()
                ),
            ),
        )
    except ArtifactError as exc:
        return _failed_report("docx", exc)
    except Exception as exc:
        return _failed_report(
            "docx",
            ArtifactError(
                "docx_reopen_failed",
                "validate",
                "The DOCX file could not be reopened.",
                details={"reason": type(exc).__name__},
            ),
        )


def validate_pptx(package: bytes) -> ValidationReport:
    """Inspect and reopen PPTX through an independent public parser."""

    try:
        inspection = inspect_ooxml(package, "pptx")
        from pptx import Presentation

        reopened = Presentation(io.BytesIO(package))
        slide_count = len(reopened.slides)
        metadata = read_model_metadata(package, "pptx")
        risks: tuple[ValidationIssue, ...] = ()
        if metadata is not None and len(metadata.slides) != slide_count:
            risks = (
                ValidationIssue(
                    error_code="pptx_metadata_slide_mismatch",
                    phase="validate",
                    message="Embedded model slide count differs from the presentation.",
                    details={
                        "metadata_slides": len(metadata.slides),
                        "actual_slides": slide_count,
                    },
                ),
            )
        status: ValidationStatus = "FAIL" if risks else "PASS"
        return ValidationReport(
            format="pptx",
            status=status,
            openable=True,
            slide_count=slide_count,
            checks=(
                "bounded_ooxml_package",
                "content_types",
                "relationships",
                "macro_free",
                "external_relationship_free",
                "python_pptx_reopen",
                *(
                    ("engine_metadata",)
                    if inspection.has_engine_metadata
                    else ()
                ),
            ),
            risks=risks,
        )
    except ArtifactError as exc:
        return _failed_report("pptx", exc)
    except Exception as exc:
        return _failed_report(
            "pptx",
            ArtifactError(
                "pptx_reopen_failed",
                "validate",
                "The PPTX file could not be reopened.",
                details={"reason": type(exc).__name__},
            ),
        )


def validate_artifact(package: bytes, artifact_format: str) -> ValidationReport:
    normalized = artifact_format.casefold().lstrip(".")
    if normalized == "docx":
        return validate_docx(package)
    if normalized == "pdf":
        from app.artifacts.pdf import validate_pdf

        return validate_pdf(package)
    if normalized == "pptx":
        return validate_pptx(package)
    return _failed_report(
        normalized,
        ArtifactError(
            "unsupported_artifact_format",
            "validate",
            "No validator is registered for this artifact format.",
            details={"format": normalized},
        ),
    )
