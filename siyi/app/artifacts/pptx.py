"""PPTX adapter backed by python-pptx 1.0.2 and the shared OOXML layer."""

from __future__ import annotations

import io
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from app.artifacts.errors import ArtifactError
from app.artifacts.model import (
    ArtifactDocument,
    ArtifactSlide,
    ArtifactTable,
    ContentBlock,
    TextReplacement,
    apply_text_replacements,
    normalize_document,
    normalize_replacements,
)
from app.artifacts.ooxml import (
    embed_model_metadata,
    inspect_ooxml,
    read_model_metadata,
    resolve_workspace_image,
)

_CJK_FONT = "Microsoft YaHei"


def _set_run_style(
    run: Any,
    *,
    size_pt: float = 18,
    bold: bool = False,
    color: str = "222222",
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Pt

    run.font.name = _CJK_FONT
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    east_asian = rpr.find(qn("a:ea"))
    if east_asian is None:
        east_asian = OxmlElement("a:ea")
        rpr.append(east_asian)
    east_asian.set("typeface", _CJK_FONT)


def _set_bullet(paragraph: Any, *, numbered: bool) -> None:
    from pptx.oxml.xmlchemy import OxmlElement

    properties = paragraph._p.get_or_add_pPr()
    for child in tuple(properties):
        if child.tag.rsplit("}", 1)[-1] in {"buNone", "buChar", "buAutoNum"}:
            properties.remove(child)
    bullet = OxmlElement("a:buAutoNum" if numbered else "a:buChar")
    if numbered:
        bullet.set("type", "arabicPeriod")
        bullet.set("startAt", "1")
    else:
        bullet.set("char", "\u2022")
    properties.insert(0, bullet)


def _add_text_box(
    slide: Any,
    text: str,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    size_pt: float,
    bold: bool = False,
) -> Any:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    shape = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    _set_run_style(run, size_pt=size_pt, bold=bold)
    return shape


def _block_height(block: ContentBlock) -> float:
    if block.kind == "heading":
        return 0.55
    if block.kind == "paragraph":
        return min(1.25, max(0.55, 0.36 * (1 + len(block.text) // 55)))
    if block.kind in {"bullet_list", "numbered_list"}:
        return max(0.55, min(2.2, 0.42 * len(block.items) + 0.12))
    if block.kind == "table" and block.table is not None:
        return min(2.7, max(0.7, 0.38 * len(block.table.rows)))
    if block.kind == "image" and block.image is not None:
        return min(3.1, max(1.2, (block.image.height_px or 240) / 120.0))
    return 0.2


def _add_list(
    slide: Any,
    block: ContentBlock,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
) -> None:
    from pptx.util import Inches

    shape = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    for index, item in enumerate(block.items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        _set_bullet(paragraph, numbered=block.kind == "numbered_list")
        paragraph.level = 0
        run = paragraph.add_run()
        run.text = item
        _set_run_style(run, size_pt=17)


def _add_table(
    slide: Any,
    table_model: ArtifactTable,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
) -> None:
    from pptx.util import Inches

    table = slide.shapes.add_table(
        len(table_model.rows),
        len(table_model.rows[0]),
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    ).table
    for row_index, row in enumerate(table_model.rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.text_frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = value
            _set_run_style(
                run,
                size_pt=13,
                bold=row_index < table_model.header_rows,
            )


def _add_image(
    slide: Any,
    block: ContentBlock,
    workspace_root: Path | None,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
) -> None:
    from pptx.util import Inches

    assert block.image is not None
    path = resolve_workspace_image(block.image.source, workspace_root)
    shape = slide.shapes.add_picture(
        str(path),
        Inches(left),
        Inches(top),
        width=Inches(width),
        height=Inches(height),
    )
    alt_text = block.image.alt_text.strip()
    if alt_text:
        shape._element.nvPicPr.cNvPr.set("descr", alt_text)
        shape._element.nvPicPr.cNvPr.set("title", alt_text[:255])


def _slides_for_document(document: ArtifactDocument) -> tuple[ArtifactSlide, ...]:
    if document.slides:
        return document.slides
    groups: list[list[ContentBlock]] = [[]]
    for block in document.blocks:
        if block.kind == "page_break":
            groups.append([])
        else:
            groups[-1].append(block)
    if not groups:
        groups = [[]]
    return tuple(
        ArtifactSlide(
            title=document.title if index == 0 else "",
            blocks=tuple(blocks),
        )
        for index, blocks in enumerate(groups)
    )


def _pptx_model(document: ArtifactDocument) -> ArtifactDocument:
    return ArtifactDocument(
        title=document.title,
        slides=_slides_for_document(document),
        metadata=dict(document.metadata),
    )


def _add_slide(
    presentation: Any,
    slide_model: ArtifactSlide,
    workspace_root: Path | None,
) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if slide_model.title:
        _add_text_box(
            slide,
            slide_model.title,
            left=0.65,
            top=0.35,
            width=12.0,
            height=0.6,
            size_pt=28,
            bold=True,
        )
    cursor = 1.15
    for block in slide_model.blocks:
        height = _block_height(block)
        if cursor + height > 7.12:
            raise ArtifactError(
                "pptx_slide_overflow",
                "create",
                "Slide content exceeds the bounded layout.",
                details={"block_kind": block.kind},
            )
        if block.kind == "heading":
            _add_text_box(
                slide,
                block.text,
                left=0.75,
                top=cursor,
                width=11.8,
                height=height,
                size_pt=max(18, 24 - (block.level or 1) * 2),
                bold=True,
            )
        elif block.kind == "paragraph":
            _add_text_box(
                slide,
                block.text,
                left=0.8,
                top=cursor,
                width=11.7,
                height=height,
                size_pt=17,
            )
        elif block.kind in {"bullet_list", "numbered_list"}:
            _add_list(
                slide,
                block,
                left=0.8,
                top=cursor,
                width=11.7,
                height=height,
            )
        elif block.kind == "table":
            assert block.table is not None
            _add_table(
                slide,
                block.table,
                left=0.8,
                top=cursor,
                width=11.7,
                height=height,
            )
        elif block.kind == "image":
            _add_image(
                slide,
                block,
                workspace_root,
                left=0.8,
                top=cursor,
                width=min(7.2, 11.7),
                height=height,
            )
        elif block.kind != "page_break":
            raise ArtifactError(
                "unsupported_artifact_block",
                "create",
                "The PPTX adapter received an unsupported block.",
                details={"kind": block.kind},
            )
        cursor += height + 0.14


def _assert_pptx_reopens(package: bytes) -> None:
    from app.artifacts.validation import validate_pptx

    report = validate_pptx(package)
    if not report.passed:
        first = report.errors[0] if report.errors else None
        details = first.details if first else {}
        if report.risks and first is None:
            details = report.risks[0].details
        raise ArtifactError(
            first.error_code if first else "pptx_reopen_failed",
            "validate",
            "The generated PPTX file could not be independently reopened.",
            details=details,
        )


def create_pptx(
    document: ArtifactDocument | Mapping[str, Any],
    *,
    workspace_root: Path | None = None,
) -> bytes:
    """Create a bounded PPTX in memory; the service layer owns filesystem writes."""

    normalized = _pptx_model(normalize_document(document))
    try:
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        presentation.core_properties.title = normalized.title
        for slide_model in normalized.slides:
            _add_slide(presentation, slide_model, workspace_root)
        buffer = io.BytesIO()
        presentation.save(buffer)
        package = embed_model_metadata(buffer.getvalue(), "pptx", normalized)
        _assert_pptx_reopens(package)
        return package
    except ArtifactError:
        raise
    except Exception as exc:
        raise ArtifactError(
            "pptx_create_failed",
            "create",
            "The PPTX file could not be created.",
            details={"reason": type(exc).__name__},
        ) from None


def _iter_shapes(shapes: Iterable[Any]) -> Iterable[Any]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _iter_shape_paragraphs(shape: Any) -> Iterable[Any]:
    if getattr(shape, "has_text_frame", False):
        yield from shape.text_frame.paragraphs
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                yield from cell.text_frame.paragraphs


def _replace_paragraph_text(paragraph: Any, old: str, new: str) -> int:
    original = "".join(run.text for run in paragraph.runs)
    count = original.count(old)
    if count == 0:
        return 0
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
    updated = "".join(run.text for run in paragraph.runs)
    if updated.count(old) == 0 and updated == original.replace(old, new):
        return count
    if not paragraph.runs:
        return 0
    paragraph.runs[0].text = original.replace(old, new)
    for run in paragraph.runs[1:]:
        run.text = ""
    return count


def _paragraph_list_kind(paragraph: Any) -> str | None:
    properties = paragraph._p.pPr
    if properties is None:
        return None
    for child in properties:
        local_name = child.tag.rsplit("}", 1)[-1]
        if local_name == "buAutoNum":
            return "numbered_list"
        if local_name == "buChar":
            return "bullet_list"
    return None


def _extract_table(shape: Any) -> ArtifactTable:
    rows = tuple(
        tuple(cell.text for cell in row.cells)
        for row in shape.table.rows
    )
    return ArtifactTable(rows=rows, header_rows=min(1, len(rows)))


def _extract_pptx_without_metadata(presentation: Any) -> ArtifactDocument:
    slides: list[ArtifactSlide] = []
    for slide in presentation.slides:
        title_shape = slide.shapes.title
        title = title_shape.text if title_shape is not None else ""
        if not title and slide.shapes:
            first = slide.shapes[0]
            if getattr(first, "has_text_frame", False):
                first_text = first.text.strip()
                if first_text:
                    title = first_text
                    title_shape = first
        blocks: list[ContentBlock] = []
        for shape in _iter_shapes(slide.shapes):
            if title_shape is not None and shape._element is title_shape._element:
                continue
            if getattr(shape, "has_table", False):
                blocks.append(ContentBlock(kind="table", table=_extract_table(shape)))
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            pending_kind: str | None = None
            pending_items: list[str] = []
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text
                if not text:
                    continue
                list_kind = _paragraph_list_kind(paragraph)
                if list_kind:
                    if pending_kind not in {None, list_kind}:
                        blocks.append(
                            ContentBlock(
                                kind=pending_kind,  # type: ignore[arg-type]
                                items=tuple(pending_items),
                            )
                        )
                        pending_items = []
                    pending_kind = list_kind
                    pending_items.append(text)
                else:
                    if pending_kind is not None:
                        blocks.append(
                            ContentBlock(
                                kind=pending_kind,  # type: ignore[arg-type]
                                items=tuple(pending_items),
                            )
                        )
                        pending_kind = None
                        pending_items = []
                    blocks.append(ContentBlock(kind="paragraph", text=text))
            if pending_kind is not None:
                blocks.append(
                    ContentBlock(
                        kind=pending_kind,  # type: ignore[arg-type]
                        items=tuple(pending_items),
                    )
                )
        slides.append(ArtifactSlide(title=title, blocks=tuple(blocks)))
    return ArtifactDocument(
        title=presentation.core_properties.title or "",
        slides=tuple(slides),
    )


def read_pptx(package: bytes) -> ArtifactDocument:
    """Securely inspect, independently reopen, and normalize a PPTX."""

    inspect_ooxml(package, "pptx")
    try:
        from pptx import Presentation

        reopened = Presentation(io.BytesIO(package))
        metadata = read_model_metadata(package, "pptx")
        return metadata if metadata is not None else _extract_pptx_without_metadata(reopened)
    except ArtifactError:
        raise
    except Exception as exc:
        raise ArtifactError(
            "pptx_read_failed",
            "read",
            "The PPTX file could not be read.",
            details={"reason": type(exc).__name__},
        ) from None


def _replace_pptx_model(
    model: ArtifactDocument,
    replacements: tuple[TextReplacement, ...],
) -> ArtifactDocument:
    visible_model = ArtifactDocument(slides=model.slides, metadata=dict(model.metadata))
    updated_visible, _ = apply_text_replacements(visible_model, replacements)
    title = model.title
    for replacement in replacements:
        title = title.replace(replacement.old, replacement.new)
    return ArtifactDocument(
        title=title,
        slides=updated_visible.slides,
        metadata=dict(model.metadata),
    )


def edit_pptx(
    package: bytes,
    replacements: (
        Mapping[str, str]
        | Sequence[TextReplacement | Mapping[str, Any]]
    ),
) -> bytes:
    """Apply count-bound exact visible-text replacements and return new bytes."""

    inspect_ooxml(package, "pptx")
    normalized_replacements = normalize_replacements(replacements)
    try:
        from pptx import Presentation

        existing_model = read_model_metadata(package, "pptx")
        presentation = Presentation(io.BytesIO(package))
        for replacement in normalized_replacements:
            actual_count = sum(
                _replace_paragraph_text(
                    paragraph,
                    replacement.old,
                    replacement.new,
                )
                for slide in presentation.slides
                for shape in _iter_shapes(slide.shapes)
                for paragraph in _iter_shape_paragraphs(shape)
            )
            if actual_count != replacement.expected_count:
                raise ArtifactError(
                    "artifact_replacement_count_mismatch",
                    "edit",
                    "Exact text replacement count did not match.",
                    details={
                        "expected_count": replacement.expected_count,
                        "actual_count": actual_count,
                    },
                )
        if existing_model is not None:
            updated_model = _replace_pptx_model(
                existing_model,
                normalized_replacements,
            )
            presentation.core_properties.title = updated_model.title
        buffer = io.BytesIO()
        presentation.save(buffer)
        raw_result = buffer.getvalue()
        if existing_model is None:
            reopened = Presentation(io.BytesIO(raw_result))
            updated_model = _extract_pptx_without_metadata(reopened)
        result = embed_model_metadata(raw_result, "pptx", updated_model)
        _assert_pptx_reopens(result)
        return result
    except ArtifactError:
        raise
    except Exception as exc:
        raise ArtifactError(
            "pptx_edit_failed",
            "edit",
            "The PPTX file could not be edited.",
            details={"reason": type(exc).__name__},
        ) from None


replace_pptx_text = edit_pptx
