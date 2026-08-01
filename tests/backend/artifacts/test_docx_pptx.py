from __future__ import annotations

import io
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from PIL import Image

from app.artifacts.docx import create_docx, edit_docx, read_docx
from app.artifacts.errors import ArtifactError
from app.artifacts.model import (
    ArtifactDocument,
    ArtifactImage,
    ArtifactSlide,
    ArtifactTable,
    ContentBlock,
)
from app.artifacts.ooxml import (
    DEFAULT_OOXML_LIMITS,
    embed_model_metadata,
    inspect_ooxml,
    read_model_metadata,
    resolve_workspace_image,
)
from app.artifacts.pptx import create_pptx, edit_pptx, read_pptx
from app.artifacts.validation import validate_docx, validate_pptx


TITLE = "\u7edf\u4e00\u4ea7\u7269\u5f15\u64ce"
OLD_TEXT = "\u7cbe\u786e\u65e7\u6587\u672c"
NEW_TEXT = "\u7cbe\u786e\u65b0\u6587\u672c"


def _document(*, with_image: bool = False) -> ArtifactDocument:
    blocks = [
        ContentBlock(kind="heading", text="\u7b2c\u4e00\u7ae0", level=1),
        ContentBlock(kind="paragraph", text=OLD_TEXT),
        ContentBlock(
            kind="bullet_list",
            items=("\u8981\u70b9\u4e00", "\u8981\u70b9\u4e8c"),
        ),
        ContentBlock(
            kind="numbered_list",
            items=("\u6b65\u9aa4\u4e00", "\u6b65\u9aa4\u4e8c"),
        ),
        ContentBlock(
            kind="table",
            table=ArtifactTable(
                rows=(
                    ("\u540d\u79f0", "\u72b6\u6001"),
                    ("\u9879\u76ee", "\u901a\u8fc7"),
                ),
            ),
        ),
    ]
    if with_image:
        blocks.append(
            ContentBlock(
                kind="image",
                image=ArtifactImage(
                    source="images/sample.png",
                    alt_text="\u5de5\u4f5c\u533a\u56fe\u7247",
                    width_px=240,
                    height_px=120,
                ),
            )
        )
    blocks.extend(
        [
            ContentBlock(kind="page_break"),
            ContentBlock(kind="paragraph", text="\u7b2c\u4e8c\u9875"),
        ]
    )
    return ArtifactDocument(
        title=TITLE,
        blocks=tuple(blocks),
        metadata={"language": "zh-CN", "engine": "siyi"},
    )


def _rewrite_package(
    package: bytes,
    *,
    replacements: dict[str, bytes] | None = None,
    removals: set[str] | None = None,
    additions: dict[str, bytes] | None = None,
) -> bytes:
    output = io.BytesIO()
    replacements = replacements or {}
    removals = removals or set()
    additions = additions or {}
    with zipfile.ZipFile(io.BytesIO(package), "r") as source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            if info.filename in removals or info.filename in additions:
                continue
            target.writestr(
                info,
                replacements.get(info.filename, source.read(info.filename)),
            )
        for name, value in additions.items():
            target.writestr(name, value)
    return output.getvalue()


def _all_docx_text(package: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(package))
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return "\n".join(parts)


def _all_pptx_text(package: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(package))
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                parts.extend(
                    cell.text
                    for row in shape.table.rows
                    for cell in row.cells
                )
    return "\n".join(parts)


def test_docx_create_rich_chinese_content_and_independent_reopen() -> None:
    package = create_docx(_document())

    report = validate_docx(package)
    assert report.status == "PASS"
    assert report.openable is True
    assert report.page_count == 2
    assert report.errors == ()
    assert TITLE in _all_docx_text(package)
    assert OLD_TEXT in _all_docx_text(package)
    assert "\u9879\u76ee" in _all_docx_text(package)

    inspection = inspect_ooxml(package, "docx")
    assert inspection.has_engine_metadata is True
    reopened = read_docx(package)
    assert reopened == _document()

    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "ListBullet" in xml
    assert "ListNumber" in xml
    assert "w:type=\"page\"" in xml
    assert _CJK_EAST_ASIA_MARKER in xml

    edited = edit_docx(
        package,
        [{"old": OLD_TEXT, "new": NEW_TEXT, "expected_count": 1}],
    )
    assert NEW_TEXT in _all_docx_text(edited)
    assert read_docx(edited).blocks[1].text == NEW_TEXT


_CJK_EAST_ASIA_MARKER = 'w:eastAsia="Microsoft YaHei"'


def test_pptx_create_rich_chinese_content_and_independent_reopen() -> None:
    package = create_pptx(_document())

    report = validate_pptx(package)
    assert report.status == "PASS"
    assert report.openable is True
    assert report.slide_count == 2
    assert TITLE in _all_pptx_text(package)
    assert OLD_TEXT in _all_pptx_text(package)
    assert "\u901a\u8fc7" in _all_pptx_text(package)
    reopened = read_pptx(package)
    assert len(reopened.slides) == 2
    assert reopened.slides[0].title == TITLE

    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")
    assert "<a:buChar" in xml
    assert "<a:buAutoNum" in xml
    assert 'typeface="Microsoft YaHei"' in xml

    edited = edit_pptx(
        package,
        [{"old": OLD_TEXT, "new": NEW_TEXT, "expected_count": 1}],
    )
    assert NEW_TEXT in _all_pptx_text(edited)
    assert read_pptx(edited).slides[0].blocks[1].text == NEW_TEXT


def test_workspace_images_are_embedded_in_docx_and_pptx(tmp_path: Path) -> None:
    image_path = tmp_path / "images" / "sample.png"
    image_path.parent.mkdir()
    Image.new("RGB", (240, 120), color=(18, 116, 142)).save(image_path)
    model = _document(with_image=True)

    docx_package = create_docx(model, workspace_root=tmp_path)
    pptx_package = create_pptx(model, workspace_root=tmp_path)

    assert inspect_ooxml(docx_package, "docx").image_count == 1
    assert inspect_ooxml(pptx_package, "pptx").image_count == 1
    assert validate_docx(docx_package).passed
    assert validate_pptx(pptx_package).passed


def test_workspace_image_path_escape_is_rejected(tmp_path: Path) -> None:
    outside = tmp_path.parent / "outside.png"
    Image.new("RGB", (2, 2)).save(outside)
    model = ArtifactDocument(
        blocks=(
            ContentBlock(
                kind="image",
                image=ArtifactImage(source="../outside.png"),
            ),
        )
    )
    with pytest.raises(ArtifactError) as exc_info:
        create_docx(model, workspace_root=tmp_path)
    assert exc_info.value.error_code == "artifact_image_path_invalid"
    assert str(tmp_path) not in str(exc_info.value.to_dict())


def test_workspace_image_format_matches_frozen_pillow_plugins(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    unsupported = workspace / "sample.gif"
    with Image.new("RGB", (8, 8), "white") as image:
        image.save(unsupported, format="GIF")

    with pytest.raises(ArtifactError) as exc_info:
        resolve_workspace_image("sample.gif", workspace_root=workspace)

    assert exc_info.value.error_code == "artifact_image_invalid"


def test_docx_exact_replacement_handles_split_runs_and_reopens() -> None:
    from docx import Document

    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("\u7cbe\u786e")
    paragraph.add_run("\u65e7\u6587\u672c")
    buffer = io.BytesIO()
    document.save(buffer)

    edited = edit_docx(
        buffer.getvalue(),
        [{"old": OLD_TEXT, "new": NEW_TEXT, "expected_count": 1}],
    )

    assert OLD_TEXT not in _all_docx_text(edited)
    assert NEW_TEXT in _all_docx_text(edited)
    assert validate_docx(edited).passed
    assert read_model_metadata(edited, "docx") is not None


def test_pptx_exact_replacement_handles_split_runs_and_reopens() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(8), Inches(1)
    ).text_frame
    paragraph = frame.paragraphs[0]
    paragraph.add_run().text = "\u7cbe\u786e"
    paragraph.add_run().text = "\u65e7\u6587\u672c"
    buffer = io.BytesIO()
    presentation.save(buffer)

    edited = edit_pptx(
        buffer.getvalue(),
        [{"old": OLD_TEXT, "new": NEW_TEXT, "expected_count": 1}],
    )

    assert OLD_TEXT not in _all_pptx_text(edited)
    assert NEW_TEXT in _all_pptx_text(edited)
    assert validate_pptx(edited).passed
    assert read_model_metadata(edited, "pptx") is not None


@pytest.mark.parametrize("factory,editor", [(create_docx, edit_docx), (create_pptx, edit_pptx)])
def test_exact_replacement_rejects_ambiguous_count(factory, editor) -> None:
    model = ArtifactDocument(
        blocks=(
            ContentBlock(kind="paragraph", text=f"{OLD_TEXT} {OLD_TEXT}"),
        )
    )
    package = factory(model)
    with pytest.raises(ArtifactError) as exc_info:
        editor(package, {OLD_TEXT: NEW_TEXT})
    assert exc_info.value.error_code == "artifact_replacement_count_mismatch"
    assert exc_info.value.safe_details == {
        "expected_count": 1,
        "actual_count": 2,
    }


@pytest.mark.parametrize(
    ("artifact_format", "factory"),
    [("docx", create_docx), ("pptx", create_pptx)],
)
def test_ooxml_rejects_zip_slip_macro_and_external_relationship(
    artifact_format: str,
    factory,
) -> None:
    package = factory(_document())
    zip_slip = _rewrite_package(
        package,
        additions={"../escape.xml": b"<escape/>"},
    )
    with pytest.raises(ArtifactError, match="traversal"):
        inspect_ooxml(zip_slip, artifact_format)

    active_part = (
        "word/vbaProject.bin"
        if artifact_format == "docx"
        else "ppt/vbaProject.bin"
    )
    macro = _rewrite_package(package, additions={active_part: b"macro"})
    with pytest.raises(ArtifactError) as macro_error:
        inspect_ooxml(macro, artifact_format)
    assert macro_error.value.error_code == "ooxml_active_content"

    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        root_relationships = archive.read("_rels/.rels")
    external_relationship = (
        b'<Relationship Id="rIdExternal" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
        b'Target="https://example.invalid/" TargetMode="External"/>'
    )
    root_relationships = root_relationships.replace(
        b"</Relationships>",
        external_relationship + b"</Relationships>",
    )
    external = _rewrite_package(
        package,
        replacements={"_rels/.rels": root_relationships},
    )
    with pytest.raises(ArtifactError) as external_error:
        inspect_ooxml(external, artifact_format)
    assert external_error.value.error_code == "ooxml_external_relationship"


@pytest.mark.parametrize(
    ("artifact_format", "factory", "main_part"),
    [
        ("docx", create_docx, "word/document.xml"),
        ("pptx", create_pptx, "ppt/presentation.xml"),
    ],
)
def test_ooxml_requires_content_types_and_main_document(
    artifact_format: str,
    factory,
    main_part: str,
) -> None:
    package = factory(_document())
    missing_types = _rewrite_package(package, removals={"[Content_Types].xml"})
    with pytest.raises(ArtifactError) as content_error:
        inspect_ooxml(missing_types, artifact_format)
    assert content_error.value.error_code == "ooxml_missing_content_types"

    missing_main = _rewrite_package(package, removals={main_part})
    with pytest.raises(ArtifactError) as main_error:
        inspect_ooxml(missing_main, artifact_format)
    assert main_error.value.error_code == "ooxml_missing_main_document"


def test_ooxml_enforces_entry_count_expansion_and_compression_ratio() -> None:
    package = create_docx(_document())
    with pytest.raises(ArtifactError) as entry_error:
        inspect_ooxml(
            package,
            "docx",
            limits=replace(DEFAULT_OOXML_LIMITS, max_entries=1),
        )
    assert entry_error.value.error_code == "ooxml_entry_limit"

    with pytest.raises(ArtifactError) as expansion_error:
        inspect_ooxml(
            package,
            "docx",
            limits=replace(DEFAULT_OOXML_LIMITS, max_uncompressed_bytes=64),
        )
    assert expansion_error.value.error_code == "ooxml_expansion_limit"

    with pytest.raises(ArtifactError) as ratio_error:
        inspect_ooxml(
            package,
            "docx",
            limits=replace(DEFAULT_OOXML_LIMITS, max_compression_ratio=1.0),
        )
    assert ratio_error.value.error_code == "ooxml_compression_ratio_limit"


def test_metadata_round_trip_is_normalized_and_replaceable() -> None:
    package = create_docx(_document())
    model = read_model_metadata(package, "docx")
    assert model == _document()

    updated = ArtifactDocument(
        title="\u66f4\u65b0\u6807\u9898",
        blocks=(ContentBlock(kind="paragraph", text="\u66f4\u65b0\u5185\u5bb9"),),
    )
    embedded = embed_model_metadata(package, "docx", updated)
    assert read_model_metadata(embedded, "docx") == updated
    assert inspect_ooxml(embedded, "docx").has_engine_metadata


def test_validation_reports_corruption_without_leaking_input() -> None:
    report = validate_docx(b"not-a-docx secret document text")
    assert report.status == "FAIL"
    assert report.openable is False
    assert report.errors[0].error_code == "ooxml_invalid_zip"
    rendered = str(report.to_dict())
    assert "secret document text" not in rendered

    error = ArtifactError(
        "safe_error",
        "test",
        details={
            "path": "C:" + r"\Users\person\private\report.docx",
            "api_key": "never-expose",
            "count": 2,
        },
    )
    assert error.safe_details["path"] == "report.docx"
    assert error.safe_details["api_key"] == "[REDACTED]"
    assert "never-expose" not in str(error.to_dict())
